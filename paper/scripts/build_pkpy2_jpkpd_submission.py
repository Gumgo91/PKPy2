"""Convert the PKPy2 manuscript build into a Journal of Pharmacokinetics and Pharmacodynamics submission.

Usage: python build_pkpy2_jpkpd_submission.py <build dir> <output dir>

<build dir> holds the manuscript produced by revise_pkpy2_peerj_manuscript.py, Tables/, the
supplement PDF and Supplemental_Files/. The Springer instructions applied here: title page with
ORCID, 4-6 keywords, unnumbered sections Abstract ... Conclusions, a Declarations section before
the references, numbered references in Springer style, tables and figures with legends after the
references ("Fig. 1" captions without final punctuation, panels a, b, c), page numbers, figure files
named Fig1.tif (600 dpi RGB) and Fig1.eps, and supplementary files named ESM_1 ... cited as Online Resources.
"""
from pathlib import Path
import copy
import re
import shutil
import sys

import docx
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BUILD = Path(sys.argv[1])
OUT = Path(sys.argv[2])
FIGS = ROOT / 'docs/pkpy2_paper/figures_peerj_revision'
TITLE = 'PKPy2: A Python framework for joint population pharmacokinetic estimation and uncertainty assessment'
KEYWORDS = ['Population pharmacokinetics', 'Pharmacokinetic-pharmacodynamic modeling', 'Nonlinear mixed-effects models',
            'Marginal likelihood', 'Software validation', 'Python']
ORCID = [('Hyunseung Kong', '0000-0001-7681-7429'), ('Inyoung Kim', '0000-0003-3062-6576')]

REFERENCES = [
    'Kong H, Kim I, Zhang B-T (2025) PKPy: a Python-based framework for automated population pharmacokinetic analysis. PeerJ 13:e20258. '
    'https://doi.org/10.7717/peerj.20258',
    'Mould DR, Upton RN (2012) Basic concepts in population modeling, simulation, and model-based drug development. CPT Pharmacometrics Syst '
    'Pharmacol 1:e6. https://doi.org/10.1038/psp.2012.4',
    'Pinheiro JC, Bates DM (1995) Approximations to the log-likelihood function in the nonlinear mixed-effects model. J Comput Graph Stat '
    '4:12–35. https://doi.org/10.1080/10618600.1995.10474663',
    'Delyon B, Lavielle M, Moulines E (1999) Convergence of a stochastic approximation version of the EM algorithm. Ann Stat 27:94–128. '
    'https://doi.org/10.1214/aos/1018031103',
    'Comets E, Lavenu A, Lavielle M (2017) Parameter estimation in nonlinear mixed effect models using saemix, an R implementation of the SAEM '
    'algorithm. J Stat Softw 80(3):1–41. https://doi.org/10.18637/jss.v080.i03',
    'Bauer RJ (2019) NONMEM tutorial part II: estimation methods and advanced examples. CPT Pharmacometrics Syst Pharmacol 8:538–556. '
    'https://doi.org/10.1002/psp4.12422',
    'Virtanen P, Gommers R, Oliphant TE et al (2020) SciPy 1.0: fundamental algorithms for scientific computing in Python. Nat Methods '
    '17:261–272. https://doi.org/10.1038/s41592-019-0686-2',
    'Lam SK, Pitrou A, Seibert S (2015) Numba: a LLVM-based Python JIT compiler. In: Proceedings of the Second Workshop on the LLVM Compiler '
    'Infrastructure in HPC. ACM, New York, Article 7. https://doi.org/10.1145/2833157.2833162',
    'Fidler ML, Wang W, Hallow M (2026) rxode2 additional model types. rxode2 documentation. '
    'https://nlmixr2.github.io/rxode2/articles/rxode2-model-types.html. Accessed 16 September 2026',
    'R Core Team (2026) Theoph: pharmacokinetics of theophylline. R datasets package documentation. '
    'https://stat.ethz.ch/R-manual/R-devel/library/datasets/html/Theoph.html. Accessed 16 September 2026',
    'Kwack H, Kong H, Lim J, Zhang B-T, Hahn J, Chang MJ (2026) PKGPT: expert-orchestrated recursive LLM agent for automated NONMEM PopPK '
    'modeling with human benchmarking. Pharmaceutics 18:501. https://doi.org/10.3390/pharmaceutics18040501',
    'Beal SL, Boeckmann AJ, Sheiner LB (1992) NONMEM users guide, part VI: PREDPP guide. NONMEM Project Group, University of California, '
    'San Francisco',
    'Beal SL, Sheiner LB (1988) NONMEM users guide, part II: users supplemental guide. NONMEM Project Group, University of California, '
    'San Francisco',
    'ICON plc (2026) NONMEM. https://www.iconplc.com/solutions/technologies/nonmem. Accessed 16 September 2026',
    'Fidler M, Wilkins JJ, Hooijmaijers R, Post TM, Schoemaker R, Trame MN, Xiong Y, Wang W (2019) Nonlinear mixed-effects model development '
    'and simulation using nlmixr and related R open-source packages. CPT Pharmacometrics Syst Pharmacol 8:621–633. '
    'https://doi.org/10.1002/psp4.12445',
    "Owen AB (1998) Scrambling Sobol' and Niederreiter–Xing points. J Complex 14:466–489. https://doi.org/10.1006/jcom.1998.0487",
    'Oehlert GW (1992) A note on the delta method. Am Stat 46:27–29. https://doi.org/10.1080/00031305.1992.10475842',
    'Golub GH, Welsch JH (1969) Calculation of Gauss quadrature rules. Math Comput 23:221–230. https://doi.org/10.1090/S0025-5718-69-99647-1',
    'Wilson EB (1927) Probable inference, the law of succession, and statistical inference. J Am Stat Assoc 22:209–212. '
    'https://doi.org/10.1080/01621459.1927.10502953',
    'Efron B (1979) Bootstrap methods: another look at the jackknife. Ann Stat 7:1–26. https://doi.org/10.1214/aos/1176344552',
    'Schoemaker R, Fidler M, Laveille C, Wilkins JJ, Hooijmaijers R, Post TM, Trame MN, Xiong Y, Wang W (2019) Performance of the SAEM and '
    'FOCEI algorithms in the open-source, nonlinear mixed effect modeling tool nlmixr. CPT Pharmacometrics Syst Pharmacol 8:923–930. '
    'https://doi.org/10.1002/psp4.12471',
    'R Core Team (2026) R: a language and environment for statistical computing, version 4.5.3. R Foundation for Statistical Computing, '
    'Vienna. https://www.R-project.org/',
    'Moler C, Van Loan C (2003) Nineteen dubious ways to compute the exponential of a matrix, twenty-five years later. SIAM Rev 45:3–49. '
    'https://doi.org/10.1137/S00361445024180',
    'Dormand JR, Prince PJ (1980) A family of embedded Runge-Kutta formulae. J Comput Appl Math 6:19–26. '
    'https://doi.org/10.1016/0771-050X(80)90013-3',
    'Shampine LF, Reichelt MW (1997) The MATLAB ODE suite. SIAM J Sci Comput 18:1–22. https://doi.org/10.1137/S1064827594276424',
    'Savic RM, Jonker DM, Kerbusch T, Karlsson MO (2007) Implementation of a transit compartment model for describing drug absorption in '
    'pharmacokinetic studies. J Pharmacokinet Pharmacodyn 34:711–726. https://doi.org/10.1007/s10928-007-9066-0',
    'Dayneka NL, Garg V, Jusko WJ (1993) Comparison of four basic models of indirect pharmacodynamic responses. J Pharmacokinet Biopharm '
    '21:457–478. https://doi.org/10.1007/BF01061691',
    'Mager DE, Jusko WJ (2001) General pharmacokinetic model for drugs exhibiting target-mediated drug disposition. J Pharmacokinet '
    'Pharmacodyn 28:507–532. https://doi.org/10.1023/A:1014414520282',
    'Gibiansky L, Gibiansky E, Kakkar T, Ma P (2008) Approximations of the target-mediated drug disposition model and identifiability of '
    'model parameters. J Pharmacokinet Pharmacodyn 35:573–591. https://doi.org/10.1007/s10928-008-9102-8',
    'Karlsson MO, Sheiner LB (1993) The importance of modeling interoccasion variability in population pharmacokinetic analyses. '
    'J Pharmacokinet Biopharm 21:735–750. https://doi.org/10.1007/BF01113502',
    'Beal SL (2001) Ways to fit a PK model with some data below the quantification limit. J Pharmacokinet Pharmacodyn 28:481–504. '
    'https://doi.org/10.1023/A:1012299115260',
    'Hooker AC, Staatz CE, Karlsson MO (2007) Conditional weighted residuals (CWRES): a model diagnostic for the FOCE method. Pharm Res '
    '24:2187–2197. https://doi.org/10.1007/s11095-007-9361-x',
    'Brendel K, Comets E, Laffont C, Laveille C, Mentré F (2006) Metrics for external model evaluation with an application to the '
    'population pharmacokinetics of gliclazide. Pharm Res 23:2036–2049. https://doi.org/10.1007/s11095-006-9067-5',
    'Comets E, Brendel K, Mentré F (2008) Computing normalised prediction distribution errors to evaluate nonlinear mixed-effect models: '
    'the npde add-on package for R. Comput Methods Programs Biomed 90:154–166. https://doi.org/10.1016/j.cmpb.2007.12.002',
    'Savic RM, Karlsson MO (2009) Importance of shrinkage in empirical Bayes estimates for diagnostics: problems and solutions. AAPS J '
    '11:558–569. https://doi.org/10.1208/s12248-009-9133-0',
    'Bergstrand M, Hooker AC, Wallin JE, Karlsson MO (2011) Prediction-corrected visual predictive checks for diagnosing nonlinear '
    'mixed-effects models. AAPS J 13:143–151. https://doi.org/10.1208/s12248-011-9255-z',
    'Dosne AG, Bergstrand M, Harling K, Karlsson MO (2016) Improving the estimation of parameter uncertainty distributions in nonlinear '
    'mixed effects models using sampling importance resampling. J Pharmacokinet Pharmacodyn 43:583–596. '
    'https://doi.org/10.1007/s10928-016-9487-8',
    'Dosne AG, Bergstrand M, Karlsson MO (2017) An automated sampling importance resampling procedure for estimating parameter '
    'uncertainty. J Pharmacokinet Pharmacodyn 44:509–520. https://doi.org/10.1007/s10928-017-9542-0',
    'Jonsson EN, Karlsson MO (1998) Automated covariate model building within NONMEM. Pharm Res 15:1463–1468. '
    'https://doi.org/10.1023/A:1011970125687',
    "O'Reilly RA, Aggeler PM (1968) Studies on coumarin anticoagulant drugs: initiation of warfarin therapy without a loading dose. "
    'Circulation 38:169–177. https://doi.org/10.1161/01.CIR.38.1.169',
]

DECLARATIONS = [
    ('Funding', 'No funding was received for conducting this study.'),
    ('Competing interests', 'The authors have no relevant financial or non-financial interests to disclose.'),
    ('Ethics approval', 'Not applicable. The study used simulated data and publicly available, de-identified pharmacokinetic datasets.'),
    ('Consent to participate', 'Not applicable.'),
    ('Consent for publication', 'Not applicable.'),
    ('Data availability', 'The theophylline data are available in the R datasets package [10], the warfarin and tobramycin data in the '
     'dataset directory of the PKGPT repository (https://github.com/Gumgo91/PKGPT) [11], and the warfarin PK/PD data in the nlmixr2data R '
     'package [15,40]. The analysis datasets, per-dataset estimates of all '
     'programs, numerical checks, and timings underlying all tables and figures are provided in Online Resources 2 and 3, with variable '
     'definitions in Online Resource 4. The nlmixr2 and saemix comparison scripts, the tobramycin analyses, and the saved fit records are '
     'available in the PKPy2 repository (https://github.com/Gumgo91/PKPy2).'),
    ('Code availability', 'PKPy2, installation instructions, runnable examples, and the numerical validation script are available at '
     'https://github.com/Gumgo91/PKPy2 under the MIT License.'),
    ('Author contributions', 'Conceptualization, methodology, software, formal analysis, investigation, data curation, visualization, and '
     'writing – original draft: Hyunseung Kong. Supervision and writing – review and editing: Inyoung Kim. All authors read and approved the '
     'final manuscript.'),
]

CAPTIONS = {
    1: 'Population inference in PKPy and PKPy2. (a) PKPy fits each subject separately and summarizes the individual log parameters. '
       '(b) PKPy2 reads event records, fits the declared population model jointly by marginal likelihood with an independent numerical '
       'audit, and provides diagnostics and interval estimates',
    2: 'Parameter recovery in the primary simulation by PKPy, the Gaussian two-stage control, and PKPy2. Left, relative bias with ±1.96 Monte '
       'Carlo standard errors; right, relative root mean squared error (RMSE). Each point summarizes 100 datasets (98 for the rich-sampling '
       'Gaussian control)',
    3: 'Coverage (a) and width (b) of PKPy2 95% confidence intervals in the primary simulation. Error bars in (a) are Wilson 95% intervals, and '
       'the dashed line marks 95%. Widths are medians relative to the true value, with interquartile ranges',
    4: 'Differences of the theophylline (a), warfarin (b), and tobramycin (c) estimates from the published expert NONMEM estimates. Tobramycin '
       'estimates were obtained with the expert-judgment constraints. Differences larger than 50% are plotted to the right of the dotted line '
       'and labeled with their values',
    5: 'Relative errors of PKPy2, nlmixr2 FOCEi, nlmixr2 SAEM, and saemix estimates in the 200 primary simulation datasets. Boxes show medians '
       'and interquartile ranges, whiskers extend to 1.5 times the interquartile range, and diamonds mark means',
    6: 'Verification of the event-record interface. (a) Maximum relative differences of PKPy2 predictions from rxode2 in linear and '
       'nonlinear (ODE) scenarios. (b) Absolute OFV differences from an independent quadrature implementation and from the compact '
       'interface; the dashed line marks the convergence tolerance of 0.05. (c) CWRES and IWRES of the theophylline model from PKPy2 and '
       'nlmixr2 at identical parameters. (d) Exact OFV at the nlmixr2 FOCEi and SAEM estimates minus that at the PKPy2 estimates',
    7: 'Warfarin PK/PD application and interval methods. (a, b) Visual predictive checks of plasma concentration and prothrombin complex '
       'activity (PCA): observed 5th, 50th, and 95th percentiles (lines), 95% intervals of the simulated percentiles (bands), and '
       'observations (points). (c) 95% intervals for the theophylline model from the Wald, sandwich, profile-likelihood, bootstrap, and SIR '
       'methods, relative to the estimate',
}

ESM = [('ESM_1.pdf', 'PKPy2_supplement.pdf', 'Supplementary methods, complete simulation summaries, clinical reference details, the '
        'comparison with nlmixr2 and saemix, the tobramycin analyses, and the methods and evaluation of the event-record interface '
        '(Sections S1-S12, Tables S1-S25)'),
       ('ESM_2.xlsx', 'Supplemental_Files/PKPy2_raw_data.xlsx', 'Raw data underlying all tables and figures: simulated and clinical analysis '
        'datasets, per-dataset estimates of all programs, numerical checks, and timings'),
       ('ESM_3.zip', 'Supplemental_Files/PKPy2_raw_data_csv.zip', 'The sheets of Online Resource 2 as CSV files'),
       ('ESM_4.xlsx', 'Supplemental_Files/PKPy2_codebook.xlsx', 'Codebook: variable definitions, units, and codes of categorical variables')]


sys.path.insert(0, str(Path(__file__).resolve().parent))
from pkpy2_references import CITATION, expand, compress, citation_order, renumber   # noqa: E402
RENUMBER = {}


def renumber_text(text):
    return renumber(text, RENUMBER)


def renumber_paragraph(p, mapping):
    """Renumber every citation of a paragraph exactly once (inside runs when no citation spans runs)."""
    text = p.text
    matches = list(CITATION.finditer(text))
    if not matches:
        return
    spans, pos = [], 0
    for r in p.runs:
        spans.append((pos, pos + len(r.text)))
        pos += len(r.text)
    if pos == len(text) and all(any(a <= m.start() and m.end() <= b for a, b in spans) for m in matches):
        for r in p.runs:
            r.text = renumber(r.text, mapping)
    else:
        set_text(p, renumber(text, mapping))

# ------------------------------------------------------------------ helpers
def find(doc, prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(hits) == 1, (prefix, len(hits))
    return hits[0]


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def sub_in_paragraph(p, pattern, repl):
    """Regex substitution inside runs; falls back to whole-paragraph text when a match spans runs."""
    if not re.search(pattern, p.text):
        return 0
    n = 0
    for r in p.runs:
        new, k = re.subn(pattern, repl, r.text)
        if k:
            r.text = new
            n += k
    if re.search(pattern, p.text):
        set_text(p, re.sub(pattern, repl, p.text))
    return n


def insert_paragraph_after(p, like=None):
    new = copy.deepcopy((like or p)._p)
    for r in new.findall(qn('w:r')):
        new.remove(r)
    p._p.addnext(new)
    return docx.text.paragraph.Paragraph(new, p._parent)


def add_runs(p, parts, size=None):
    """parts: list of (text, bold)."""
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = 'Times New Roman'
        if size:
            r.font.size = Pt(size)
    return p


def remove(p):
    p._p.getparent().remove(p._p)


def page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = 1
    run = p.add_run()
    for tag, text in [('begin', None), (None, 'PAGE'), ('end', None)]:
        if tag:
            el = OxmlElement('w:fldChar')
            el.set(qn('w:fldCharType'), tag)
        else:
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = text
        run._r.append(el)


def new_section(doc):
    """New section on a new page that continues the page numbering."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    numbering = section._sectPr.find(qn('w:pgNumType'))
    if numbering is not None and numbering.get(qn('w:start')) is not None:
        del numbering.attrib[qn('w:start')]
    return section


def compact(tbl):
    """Single-spaced cell paragraphs without extra space so that each table fits on one page."""
    for p in tbl.iter(qn('w:p')):
        ppr = p.find(qn('w:pPr'))
        if ppr is None:
            ppr = OxmlElement('w:pPr')
            p.insert(0, ppr)
        spacing = ppr.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            ppr.append(spacing)
        for key, value in [('w:before', '0'), ('w:after', '0'), ('w:line', '240'), ('w:lineRule', 'auto')]:
            spacing.set(qn(key), value)
    for mar in tbl.iter(qn('w:tcMar')):
        for side in ('w:top', 'w:bottom'):
            el = mar.find(qn(side))
            if el is not None:
                el.set(qn('w:w'), '30')


def body_insert(doc, element):
    body = doc.element.body
    body.insert(len(body) - 1, element)


# ------------------------------------------------------------------ manuscript
def manuscript():
    d = docx.Document(BUILD / 'PKPy2_PeerJ_manuscript_revised.docx')
    normal = [p for p in d.paragraphs if p.text.startswith('Population pharmacokinetic (PopPK) models quantify')][0]

    # Title page
    for p in d.paragraphs[:12]:
        if 'Republic of South Korea' in p.text:
            sub_in_paragraph(p, 'Republic of South Korea', 'Republic of Korea')
    email = find(d, 'Email address: inyoungkim@korea.kr')
    orcid = insert_paragraph_after(email)
    add_runs(orcid, [('ORCID: ' + '; '.join(f'{n}, https://orcid.org/{o}' for n, o in ORCID), False)])

    # Keywords after the abstract
    abstract = find(d, 'Population pharmacokinetic (PopPK) analysis usually relies')
    assert 150 <= len(abstract.text.split()) <= 250
    kw = insert_paragraph_after(abstract)
    add_runs(kw, [('Keywords ', True), (' · '.join(KEYWORDS), False)])

    # Section names and cross-references
    set_text(find(d, 'Materials & Methods'), 'Methods')
    n_fig = 0
    for p in d.paragraphs:
        # In-text citations end with ')'; a sentence-initial "Figure 1" is kept in full.
        n_fig += sub_in_paragraph(p, r'Figure (\d)([A-Z]?)(?=\))', lambda m: f'Fig. {m.group(1)}{m.group(2).lower()}')
        sub_in_paragraph(p, r'Supplementary Tables (S\d+) and (S\d+)', r'Online Resource 1, Tables \1 and \2')
        sub_in_paragraph(p, r'Supplementary (Section|Table) (S\d+)', r'Online Resource 1, \1 \2')
        sub_in_paragraph(p, r'the Supplementary Material', 'Online Resource 1')
    assert n_fig >= 7, n_fig
    assert not any('Supplementary' in p.text for p in d.paragraphs)

    # Replace Data availability and Acknowledgements with Declarations before the references
    head = find(d, 'Data and code availability')
    refs_head = find(d, 'References')
    body = d.element.body
    children = list(body)
    start, stop = children.index(head._p), children.index(refs_head._p)
    for el in children[start + 1:stop]:
        body.remove(el)
    set_text(head, 'Declarations')
    anchor = head
    declarations = list(DECLARATIONS)
    for name, text in declarations:
        anchor = insert_paragraph_after(anchor, like=normal)
        add_runs(anchor, [(name + ' ', True), (text, False)])
        if 'XXXXXXX' in text:
            anchor.runs[-1].font.highlight_color = WD_COLOR_INDEX.YELLOW
    before = children[start - 1]
    if before.tag == qn('w:p') and not ''.join(t.text or '' for t in before.iter(qn('w:t'))).strip():
        anchor._p.addnext(copy.deepcopy(before))

    # References in Springer style, numbered in order of first citation
    refs = [p for p in d.paragraphs if re.match(r'\[\d+\] ', p.text)]
    assert len(refs) == len(REFERENCES)
    for i, p in enumerate(refs, start=1):
        assert p.text.startswith(f'[{i}] ')
    body_paras = [p for p in d.paragraphs if p not in refs]
    new_number = citation_order([p.text for p in body_paras], len(REFERENCES))
    order = sorted(new_number, key=new_number.get)
    for p in body_paras:
        renumber_paragraph(p, new_number)
    for i, (p, old_k) in enumerate(zip(refs, order), start=1):
        set_text(p, f'{i}. {REFERENCES[old_k - 1]}')
    RENUMBER.update(new_number)

    # Tables with legends (landscape section), then figures with legends (portrait section)
    first = d.sections[0]
    page_number_footer(first)
    width, height = first.page_width, first.page_height     # portrait; `first` tracks the last section after add_section
    first_margin = first.top_margin
    tables = new_section(d)
    tables.orientation = WD_ORIENT.LANDSCAPE
    tables.page_width, tables.page_height = height, width
    tables.top_margin = tables.bottom_margin = tables.left_margin = tables.right_margin = Inches(0.7)
    usable = int((tables.page_width - tables.left_margin - tables.right_margin) / 635)   # EMU -> twips
    for n in range(1, 7):
        src = docx.Document(BUILD / f'Tables/Table_{n}.docx')
        title, foot = src.paragraphs[0].text, [p.text for p in src.paragraphs[1:] if p.text.strip()]
        m = re.match(rf'Table {n}\. (.*?)\.?$', title)
        cap = d.add_paragraph()
        if n > 1:
            cap.paragraph_format.page_break_before = True
        add_runs(cap, [(f'Table {n}', True), (' ' + m.group(1), False)], size=10)
        tbl = copy.deepcopy(src.tables[0]._tbl)
        total = sum(int(g.get(qn('w:w'))) for g in tbl.find(qn('w:tblGrid')))
        scale = usable / total
        for el in tbl.iter():
            if el.tag in (qn('w:gridCol'), qn('w:tcW'), qn('w:tblW')) and el.get(qn('w:w')) and el.get(qn('w:type'), 'dxa') == 'dxa':
                el.set(qn('w:w'), str(int(int(el.get(qn('w:w'))) * scale)))
        compact(tbl)
        body_insert(d, tbl)
        for text in foot:
            text = renumber_text(text)
            text = re.sub(r'Supplementary Tables? (S\d+)', r'Online Resource 1, Table \1', text)
            text = re.sub(r'Supplementary (S\d+)', r'Online Resource 1, Section \1', text)
            assert 'Supplementary' not in text, text
            add_runs(d.add_paragraph(), [(text, False)], size=9)

    figs = new_section(d)
    figs.orientation = WD_ORIENT.PORTRAIT
    figs.page_width, figs.page_height = width, height
    figs.top_margin = figs.bottom_margin = first_margin
    tmp = OUT / '_embed'
    tmp.mkdir(parents=True, exist_ok=True)
    for n in range(1, 8):
        img = Image.open(FIGS / f'Figure_{n}.png').convert('RGB')
        w_in, h_in = img.size[0] / 600, img.size[1] / 600
        img = img.resize((int(img.size[0] / 2), int(img.size[1] / 2)), Image.LANCZOS)
        path = tmp / f'Fig{n}.png'
        img.save(path, dpi=(300, 300))
        scale = min(6.5 / w_in, 8.2 / h_in, 1.0)
        pic = d.add_paragraph()
        if n > 1:
            pic.paragraph_format.page_break_before = True
        pic.add_run().add_picture(str(path), width=Inches(w_in * scale))
        add_runs(d.add_paragraph(), [(f'Fig. {n}', True), (' ' + CAPTIONS[n], False)], size=10)
    d.save(OUT / 'PKPy2_JPKPD_manuscript.docx')
    shutil.rmtree(tmp)
    return d


def files():
    fig_dir = OUT / 'Figures'
    fig_dir.mkdir(parents=True, exist_ok=True)
    for n in range(1, 8):
        shutil.copy2(FIGS / f'Figure_{n}.tif', fig_dir / f'Fig{n}.tif')
        shutil.copy2(FIGS / f'Figure_{n}.eps', fig_dir / f'Fig{n}.eps')
    esm_dir = OUT / 'ESM'
    esm_dir.mkdir(parents=True, exist_ok=True)
    for name, src, _ in ESM:
        shutil.copy2(BUILD / src, esm_dir / name)
    table_dir = OUT / 'Tables_editable'
    table_dir.mkdir(parents=True, exist_ok=True)
    for n in range(1, 7):
        shutil.copy2(BUILD / f'Tables/Table_{n}.docx', table_dir / f'Table{n}.docx')
    (OUT / 'ESM_captions.txt').write_text('\n'.join(f'Online Resource {i}: {name} - {cap}' for i, (name, _, cap) in enumerate(ESM, start=1)),
                                          encoding='utf-8')


def cover_letter():
    d = docx.Document()
    style = d.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    paras = [
        'Dear Editor,',
        f'We submit the Original Paper "{TITLE}" for consideration in the Journal of Pharmacokinetics and Pharmacodynamics.',
        'PKPy2 is a standalone Python package for population pharmacokinetic and pharmacodynamic analysis that estimates structural '
        'parameters, interindividual variability, residual error, and covariate effects jointly by marginal likelihood, with explicit fixed '
        'values and parameter bounds, independent numerical convergence checks, model diagnostics, and interval estimates. It reads '
        'NONMEM-format event records and supports multi-compartment, nonlinear, and PK/PD models with correlated and interoccasion random '
        'effects, time-varying covariates, and censored observations. We verified its predictions, likelihoods, and diagnostics against '
        'rxode2, independent quadrature, nlmixr2, and the npde package, evaluated parameter recovery and interval coverage in simulated '
        'datasets, compared it with nlmixr2 (FOCEi and SAEM) and saemix on identical data, and applied it to the theophylline, warfarin, '
        'and tobramycin datasets with published expert NONMEM analyses and to a warfarin PK/PD model. The tobramycin analysis shows how '
        'declared fixed values and bounds let a pharmacometrician steer an analysis whose data support a different maximum-likelihood '
        'solution, which we expect to be of interest to the readers of the journal.',
        'PKPy2 extends our earlier PKPy framework (Kong et al., PeerJ 2025), whose estimation components serve as a comparator, and the expert '
        'NONMEM reference estimates are those published in the PKGPT study (Kwack et al., Pharmaceutics 2026), of which the first author is a '
        'co-author. No text or figures are reused from these articles. The manuscript has not been published and is not under consideration '
        'elsewhere, and all authors approved its submission. Raw data are provided as Online Resources, and the software, analysis scripts, '
        'and fit records are available in the PKPy2 GitHub repository.',
        'Thank you for considering our manuscript.',
        'Sincerely,',
        'Inyoung Kim (corresponding author)\nDepartment of Defense Science, Korea National Defense University, Nonsan, Republic of Korea\n'
        'inyoungkim@korea.kr',
    ]
    for text in paras:
        d.add_paragraph(text)
    d.save(OUT / 'Cover_letter_draft.docx')


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    manuscript()
    files()
    cover_letter()
    print('written to', OUT)
